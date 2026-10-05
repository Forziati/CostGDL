import base64
import io
import json
import re

import anthropic
import streamlit as st

from kernel import SYSTEM_PROMPT

MODELS = {
    "Sonnet 5.5 (mejor calidad)": {
        "id": "claude-sonnet-5-5", "search": "web_search_20260209", "in": 2.0, "out": 10.0,
    },
    "Haiku 4.5 (económico, para probar)": {
        "id": "claude-haiku-4-5-20251001", "search": "web_search_20250305", "in": 1.0, "out": 5.0,
    },
}
SEARCH_USD = 0.01  # aprox. por búsqueda web (USD 10 por 1000)
MAX_SEARCHES = 6
MAX_ROUNDS = 4
FILE_MAX_SEARCHES = 12  # cuando se cotiza un listado subido
FILE_MAX_ROUNDS = 6
MAX_FILE_MB = 10
MAX_TEXT_CHARS = 60_000


def web_search_tool(search_type, max_uses=MAX_SEARCHES):
    return {
        "type": search_type,
        "name": "web_search",
        "max_uses": max_uses,
        "user_location": {
            "type": "approximate",
            "city": "Guadalajara",
            "region": "Jalisco",
            "country": "MX",
            "timezone": "America/Mexico_City",
        },
    }


st.set_page_config(page_title="CostGDL", page_icon="🏗️", layout="wide")
st.title("🏗️ CostGDL · Presupuesto de remodelación")

# --- acceso opcional por contraseña -------------------------------------
pwd = st.secrets.get("APP_PASSWORD")
if pwd and st.text_input("Contraseña", type="password") != pwd:
    st.stop()

if "ANTHROPIC_API_KEY" not in st.secrets:
    st.error("Falta ANTHROPIC_API_KEY en .streamlit/secrets.toml (o en Secrets de Streamlit Cloud).")
    st.stop()

client = anthropic.Anthropic(api_key=st.secrets["ANTHROPIC_API_KEY"])

@st.cache_resource
def ledger():
    """Gasto acumulado estimado (USD). Sobrevive a recargas, no a reinicios de la app."""
    return {"spent": 0.0}


# api_messages: historial para la API (solo texto; los resultados de búsqueda no se reenvían)
# shown: solo texto, para mostrar; versions: cada respuesta del agente
st.session_state.setdefault("api_messages", [])
st.session_state.setdefault("shown", [])
st.session_state.setdefault("model_label", "Haiku 4.5 (económico, para probar)")


def run_agent(live, with_file=False):
    """Ejecuta un turno del agente (con streaming y búsqueda web).

    Dentro del turno se reenvían los bloques de búsqueda para poder reanudar un
    `pause_turn`; al final solo se guarda en el historial el texto de la respuesta,
    para no pagar en cada turno los resultados de búsqueda ya consumidos.
    Devuelve (texto, info) con info = búsquedas, tokens, costo aprox. y diagnóstico.
    """
    cfg = MODELS[st.session_state.model_label]
    turn = list(st.session_state.api_messages)
    streamed, final_text, notice = "", [], ""
    info = {"searches": 0, "in": 0, "out": 0, "rounds": [], "errors": []}
    rounds = FILE_MAX_ROUNDS if with_file else MAX_ROUNDS
    searches = FILE_MAX_SEARCHES if with_file else MAX_SEARCHES
    for _ in range(rounds):
        with client.messages.stream(
            model=cfg["id"],
            max_tokens=16000,
            system=SYSTEM_PROMPT,
            tools=[web_search_tool(cfg["search"], searches)],
            messages=turn,
        ) as stream:
            for chunk in stream.text_stream:
                streamed += chunk
                live.markdown(streamed + "▌")
            resp = stream.get_final_message()
        turn.append({"role": "assistant", "content": resp.content})
        info["in"] += resp.usage.input_tokens
        info["out"] += resp.usage.output_tokens
        info["rounds"].append(
            f"{resp.stop_reason} · {[b.type for b in resp.content]} · "
            f"in={resp.usage.input_tokens} out={resp.usage.output_tokens}"
        )
        for b in resp.content:
            if b.type == "text":
                final_text.append(b.text)
            else:
                final_text = []
                info["searches"] += b.type == "server_tool_use" and getattr(b, "name", "") == "web_search"
                content = getattr(b, "content", None)
                if b.type == "web_search_tool_result" and not isinstance(content, list):
                    info["errors"].append(str(getattr(content, "error_code", content)))
        if resp.stop_reason == "pause_turn":
            streamed += "\n\n"
            continue
        if resp.stop_reason == "max_tokens":
            notice = "\n\n⚠️ La respuesta se cortó por longitud. Escribe «continúa» para completarla."
        break
    else:
        notice = "\n\n⚠️ Se alcanzó el máximo de rondas de búsqueda. Escribe «continúa» para terminar la propuesta."
    text = "".join(final_text).strip() or streamed.strip()
    st.session_state.api_messages.append({"role": "assistant", "content": text})
    info["usd"] = (
        info["in"] * cfg["in"] + info["out"] * cfg["out"]
    ) / 1e6 + info["searches"] * SEARCH_USD
    return text + notice, info


FORM_RE = re.compile(r"```cuestionario\s*(\{.*?\})\s*```", re.S)
OTRA = "Otra (especificar abajo)"


def split_form(text):
    """Separa el texto visible del cuestionario JSON (None si no hay o es inválido)."""
    m = FORM_RE.search(text)
    if not m:
        return text, None
    try:
        form = json.loads(m.group(1))["preguntas"]
    except (ValueError, KeyError, TypeError):
        return text, None
    return (text[: m.start()] + text[m.end():]).strip(), form


def render_form(msg_idx, questions):
    """Dibuja el cuestionario; devuelve el texto de respuestas al enviarse, o None."""
    with st.form(f"form_{msg_idx}"):
        picked = {}
        for q in questions:
            opts = list(q["opciones"]) + [OTRA]
            default = q.get("defecto")
            st.markdown(f"**{q['texto']}**")
            if q.get("tipo") == "varias":
                sel = st.multiselect(
                    q["texto"], opts, default=[d for d in (default or []) if d in opts],
                    key=f"{msg_idx}_{q['id']}", label_visibility="collapsed",
                )
            else:
                idx = opts.index(default) if default in opts else 0
                sel = st.radio(
                    q["texto"], opts, index=idx, key=f"{msg_idx}_{q['id']}",
                    label_visibility="collapsed",
                )
            other = st.text_input(
                "Otra:", key=f"{msg_idx}_{q['id']}_otra", placeholder="Si elegiste «Otra», escribe aquí",
                label_visibility="collapsed",
            )
            picked[q["id"]] = (q["texto"], sel, other)
        c1, c2 = st.columns(2)
        send = c1.form_submit_button("Enviar respuestas", type="primary")
        assume = c2.form_submit_button("Asumir todo (valores por defecto)")
    if assume:
        return "Asume todos los valores por defecto del cuestionario y arma la propuesta base."
    if not send:
        return None
    lines = ["Respuestas al cuestionario:"]
    for _id, (texto, sel, other) in picked.items():
        sel = sel if isinstance(sel, list) else [sel]
        vals = [v for v in sel if v != OTRA]
        if OTRA in sel and other.strip():
            vals.append(other.strip())
        lines.append(f"- {texto} {', '.join(vals) if vals else '(sin preferencia, asume)'}")
    return "\n".join(lines)


with st.sidebar:
    st.header("Crédito")
    start = st.number_input(
        "Saldo inicial (USD)",
        min_value=0.0,
        value=float(st.secrets.get("CREDIT_START", 4.21)),
        step=0.5,
        help="Mira tu saldo en console.anthropic.com → Facturación y anótalo aquí.",
    )
    left = max(start - ledger()["spent"], 0.0)
    st.metric("Saldo estimado", f"USD {left:.2f}", f"-{ledger()['spent']:.2f} gastado", delta_color="off")
    st.progress(min(left / start, 1.0) if start else 0.0)
    if left < 0.5:
        st.warning("Saldo estimado bajo: carga crédito en la consola.")
    if st.button("Reiniciar contador"):
        ledger()["spent"] = 0.0
        st.rerun()
    st.caption("Estimación según tokens y búsquedas; el saldo oficial está en la consola.")
    st.header("Modelo")
    st.radio("Modelo", list(MODELS), key="model_label", label_visibility="collapsed")
    st.header("Versiones")
    versions = [
        split_form(m["content"])[0]
        for m in st.session_state.shown
        if m["role"] == "assistant" and split_form(m["content"])[1] is None
    ]
    if versions:
        st.download_button(
            "⬇️ Descargar última versión (.md)",
            versions[-1],
            file_name=f"presupuesto_v{len(versions)}.md",
        )
        st.caption(f"{len(versions)} versión(es) en esta sesión")
    if st.button("🗑️ Nuevo presupuesto"):
        st.session_state.api_messages = []
        st.session_state.shown = []
        st.rerun()

def file_to_blocks(f):
    """Convierte un archivo subido en bloques de contenido para la API."""
    name, data = f.name, f.getvalue()
    ext = name.rsplit(".", 1)[-1].lower()
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ValueError(f"{name}: pesa más de {MAX_FILE_MB} MB.")
    b64 = lambda: base64.standard_b64encode(data).decode()
    if ext == "pdf":
        return [{"type": "document", "title": name,
                 "source": {"type": "base64", "media_type": "application/pdf", "data": b64()}}]
    if ext in ("png", "jpg", "jpeg", "webp", "gif"):
        mt = {"jpg": "image/jpeg"}.get(ext, f"image/{ext}")
        return [{"type": "image", "source": {"type": "base64", "media_type": mt, "data": b64()}}]
    if ext in ("txt", "md", "csv"):
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
    elif ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            rows = [" | ".join("" if c is None else str(c) for c in r) for r in ws.iter_rows(values_only=True)]
            parts.append(f"## Hoja: {ws.title}\n" + "\n".join(r for r in rows if r.strip(" |")))
        text = "\n\n".join(parts)
    elif ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            lines += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        text = "\n".join(lines)
    else:
        raise ValueError(f"{name}: formato no soportado.")
    if len(text) > MAX_TEXT_CHARS:
        raise ValueError(f"{name}: es demasiado largo ({len(text):,} caracteres). Divídelo en partes.")
    return [{"type": "text", "text": f"Archivo «{name}»:\n{text}"}]


def handle_message(prompt, blocks=None, files_note=""):
    st.session_state.shown.append({"role": "user", "content": prompt + files_note})
    st.session_state.api_messages.append(
        {"role": "user", "content": (blocks or []) + [{"type": "text", "text": prompt}]}
        if blocks
        else {"role": "user", "content": prompt}
    )
    with st.chat_message("user"):
        st.markdown(prompt + files_note)
    with st.chat_message("assistant"):
        live = st.empty()
        live.markdown("Pensando… 🔎")
        try:
            answer, info = run_agent(live, with_file=bool(blocks))
        except anthropic.APIError as e:
            st.error(f"Error de la API: {e}")
            st.session_state.shown.pop()
            st.session_state.api_messages.pop()
            st.stop()
        live.markdown(split_form(answer)[0])
        ledger()["spent"] += info["usd"]
    if blocks:  # el listado ya quedó reflejado en la propuesta; no se reenvía en cada turno
        st.session_state.api_messages[-2] = {
            "role": "user", "content": prompt + files_note + " (archivo ya procesado)",
        }
    st.session_state.shown.append({"role": "assistant", "content": answer, "info": info})
    st.rerun()


last = len(st.session_state.shown) - 1
pending = None
for i, m in enumerate(st.session_state.shown):
    with st.chat_message(m["role"]):
        if m["role"] == "user":
            st.markdown(m["content"])
            continue
        text, form = split_form(m["content"])
        st.markdown(text)
        if form and i == last:
            pending = render_form(i, form)
        info = m.get("info")
        if info:
            st.caption(
                f"{info['searches']} búsqueda(s) · {info['in']:,} tokens entrada · "
                f"{info['out']:,} salida · ≈ USD {info['usd']:.2f}"
            )
            with st.expander("Diagnóstico"):
                st.write(info["rounds"])
                if info["errors"]:
                    st.warning(f"Errores de búsqueda: {info['errors']}")

if pending:
    handle_message(pending)

st.session_state.setdefault("up_n", 0)
with st.expander("📎 Subir listado de materiales (Excel, CSV, PDF, Word, TXT o foto)"):
    files = st.file_uploader(
        "Archivos", type=["xlsx", "csv", "pdf", "docx", "txt", "md", "png", "jpg", "jpeg", "webp"],
        accept_multiple_files=True, key=f"up_{st.session_state.up_n}", label_visibility="collapsed",
    )
    note = st.text_input("Instrucciones (opcional)", placeholder="Ej: es la cocina; usa porcelanato de 60x60 como alternativa")
    if st.button("Cotizar archivo(s)", type="primary", disabled=not files):
        try:
            blocks = [b for f in files for b in file_to_blocks(f)]
        except Exception as e:  # archivo ilegible o demasiado grande
            st.error(str(e))
            st.stop()
        st.session_state.up_n += 1  # limpia el cargador
        handle_message(
            "Cotiza este listado de materiales: busca precios vigentes por renglón, "
            "respeta cantidades y unidades, y arma la propuesta con subtotales y total. "
            + note,
            blocks=blocks,
            files_note="\n\n📎 " + ", ".join(f.name for f in files),
        )
if prompt := st.chat_input("Ej: cocina 4x3 m con isla y baño de 2x2 m, estilo moderno…"):
    handle_message(prompt)
